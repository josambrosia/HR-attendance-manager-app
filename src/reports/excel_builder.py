from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from src.core.constants import COLORS, REASON_CODES

PRIMARY_FILL = PatternFill("solid", fgColor=COLORS["primary"].lstrip("#"))
HEADER_FONT = Font(bold=True, color="FFFFFF", size=11)
THIN_BORDER = Border(
    left=Side(style="thin", color="CCCCCC"),
    right=Side(style="thin", color="CCCCCC"),
    top=Side(style="thin", color="CCCCCC"),
    bottom=Side(style="thin", color="CCCCCC"),
)


def build_raw_excel(output_path: str, sections: dict):
    """sections: {sheet_name: list_of_dicts}. Each sheet gets its own worksheet."""
    wb = Workbook()
    wb.remove(wb.active)
    for sheet_name, rows in sections.items():
        ws = wb.create_sheet(title=sheet_name[:31])  # Excel sheet name limit
        if not rows:
            ws.append(["(no data)"])
            continue
        headers = list(rows[0].keys())
        ws.append(headers)
        for cell in ws[1]:
            cell.font = HEADER_FONT
            cell.fill = PRIMARY_FILL
            cell.alignment = Alignment(horizontal="center")
            cell.border = THIN_BORDER
        for row in rows:
            ws.append([row.get(h, "") for h in headers])
        for col_idx, h in enumerate(headers, 1):
            max_len = max([len(str(h))] + [len(str(r.get(h, ""))) for r in rows])
            ws.column_dimensions[get_column_letter(col_idx)].width = min(max_len + 2, 40)
    wb.save(output_path)


def build_hall_of_late_excel(output_path: str, ranking: list[dict]):
    """Convenience wrapper: ranking -> single-sheet excel."""
    rows = [{
        "Rank": i + 1,
        "Staff No": r["staff_no"],
        "Name": r["name"],
        "Total Late (min)": r["total_late"],
        "Days Late": r["days_late"],
        "Max Late (min)": r["max_late"],
        "Max Late Date": r["max_late_date"],
    } for i, r in enumerate(ranking)]
    build_raw_excel(output_path, {"Hall of Late": rows})


def build_monthly_sheets_format(output_path: str, monthly_rows: list[dict]):
    """Write monthly recap matching the user's Google Sheets format (17 cols).

    Each row dict expects these keys (others are tolerated and ignored):
        name, department, date (ISO), day_name, day_type,
        schedule_in, schedule_out, actual_in, actual_out,
        work_hours, overtime_hours, kurang_hours,
        late_minutes, early_leave_minutes,
        absent_flag, forgot_punch_flag, ijin_flag,
        alasan_ijin (free-form text)

    Saturday/Sunday rows (day_type='Istirahat') write only Nama/Dept/Tanggal/Hari/Tipe
    (other columns left blank), matching the user's reference file.
    """
    headers = [
        "Nama", "Dept.", "Tanggal", "Hari", "Tipe",
        "Jadwal", "Masuk", "Keluar",
        "Kerja", "Lembur", "Kurang", "Terlambat", "Pulang Cepat",
        "Absen", "Lupa in/out", "Ijin", "Alasan Ijin",
    ]
    wb = Workbook()
    wb.remove(wb.active)
    ws = wb.create_sheet(title="Monthly Report")
    ws.append(headers)
    for cell in ws[1]:
        cell.font = HEADER_FONT
        cell.fill = PRIMARY_FILL
        cell.alignment = Alignment(horizontal="center")
        cell.border = THIN_BORDER

    for r in monthly_rows:
        is_rest = r.get("day_type") == "Istirahat"
        if is_rest:
            ws.append([
                r.get("name", ""),
                r.get("department", ""),
                r.get("date", ""),
                r.get("day_name", ""),
                r.get("day_type", "Istirahat"),
                "", "", "", "", "", "", "", "", "", "", "", "",
            ])
        else:
            schedule = ""
            if r.get("schedule_in") and r.get("schedule_out"):
                schedule = f"{r['schedule_in'].replace(':','.')} - {r['schedule_out'].replace(':','.')}"
            ws.append([
                r.get("name", ""),
                r.get("department", ""),
                r.get("date", ""),
                r.get("day_name", ""),
                r.get("day_type", ""),
                schedule,
                r.get("actual_in") or "",
                r.get("actual_out") or "",
                r.get("work_hours") if r.get("work_hours") is not None else "",
                r.get("overtime_hours") if r.get("overtime_hours") not in (None, 0) else "",
                r.get("kurang_hours") if r.get("kurang_hours") is not None else "",
                r.get("late_minutes") if r.get("late_minutes") not in (None, 0) else "",
                r.get("early_leave_minutes") if r.get("early_leave_minutes") not in (None, 0) else "",
                r.get("absent_flag") if r.get("absent_flag") not in (None, 0) else "",
                r.get("forgot_punch_flag") if r.get("forgot_punch_flag") not in (None, 0) else "",
                r.get("ijin_flag") if r.get("ijin_flag") not in (None, 0) else "",
                r.get("alasan_ijin", ""),
            ])

    # Auto-width columns
    for col_idx in range(1, 18):
        col_letter = get_column_letter(col_idx)
        max_len = max(
            [len(str(headers[col_idx - 1]))] +
            [len(str(ws.cell(row_idx, col_idx).value or ""))
             for row_idx in range(2, ws.max_row + 1)]
        )
        ws.column_dimensions[col_letter].width = min(max_len + 2, 40)

    wb.save(output_path)
